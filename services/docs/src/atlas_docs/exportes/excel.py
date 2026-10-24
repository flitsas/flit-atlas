"""Excel (openpyxl) con la forma del Excel de contratos que ya conocen los clientes (HU 4.5).

Imita la estructura del Excel de contratos ICT, no sus datos: una hoja «Índice» con la versión, la fecha y el listado
de servicios, y una hoja por servicio con «Ambiente URL» (una fila por ambiente), versión y fecha a la derecha,
«Header», «Método», «Parámetros» (Campo, Descripción, Tipo de Dato, Obligatoriedad u «OBLIGATORIO <tipo de trámite>»
cuando hay matriz, Observaciones), «Ejemplo», «Salida» y los códigos de respuesta; las tablas de referencia van a la
derecha, como en el original. Encabezados en Arial negrita blanca sobre el morado del Excel, bordes finos y texto
ajustado. Se agregan una «Portada», «Cambios» y «Tablas de referencia» (portada, changelog y marca de agua).
"""

import io
import re
from collections.abc import Sequence
from datetime import datetime, time
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from atlas_docs.exportes.contenido import SIGLAS_OBLIGATORIEDAD, Campo, Contrato, Matriz, Operacion, Tabla
from atlas_docs.exportes.documento import Documento
from atlas_docs.exportes.empaquetar import normalizar_zip

MORADO = "FF4E00EB"  # el de los encabezados del Excel de contratos
_FINO = Side(style="thin", color="FF000000")
BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
ENCABEZADO = Font(name="Arial", size=11, bold=True, color="FFFFFFFF")
RELLENO = PatternFill(fill_type="solid", fgColor=MORADO)
ETIQUETA = Font(name="Arial", size=11, bold=True, color="FF000000")
VALOR = Font(name="Calibri", size=11, color="FF000000")
CAMPO = Font(name="Calibri", size=11, bold=True, color="FF000000")
CODIGO = Font(name="Consolas", size=10, color="FF000000")
AJUSTE = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(wrap_text=True, vertical="center")

HOJAS_FIJAS = ("Portada", "Índice", "Cambios", "Tablas de referencia", "Guías")
_PROHIBIDOS = re.compile(r"[\[\]:*?/\\]")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MAX_CELDA = 32_000  # Excel admite 32.767 caracteres por celda


def _limpio(valor: object) -> str:
    texto = _CONTROL.sub("", "" if valor is None else str(valor))
    return texto if len(texto) <= MAX_CELDA else texto[: MAX_CELDA - 1] + "…"


def _como_texto(celda: Cell) -> Cell:
    """Fija la celda como texto (F4-SEC-11).

    openpyxl guarda como fórmula todo valor que empieza por «=», y el contenido sale de specs y guías que escribe otra
    gente: con el tipo «s» Excel lo muestra tal cual y nunca lo evalúa («+», «-», «@», tabulador o retorno de carro
    al inicio tampoco se interpretan en una celda de texto).
    """
    if isinstance(celda.value, str):
        celda.data_type = "s"
    return celda


def _celda(hoja: Worksheet, fila: int, col: int, valor: object, fuente: Font = VALOR, borde: bool = True) -> None:
    celda = _como_texto(hoja.cell(fila, col, _limpio(valor)))
    celda.font = fuente
    celda.alignment = AJUSTE
    if borde:
        celda.border = BORDE


def _encabezado(hoja: Worksheet, fila: int, col: int, texto: str, hasta: int | None = None) -> None:
    for c in range(col, (hasta or col) + 1):
        celda = hoja.cell(fila, c)
        celda.fill = RELLENO
        celda.border = BORDE
    celda = _como_texto(hoja.cell(fila, col, _limpio(texto)))
    celda.font = ENCABEZADO
    celda.alignment = CENTRO
    if hasta and hasta > col:
        hoja.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=hasta)


def _bloque(hoja: Worksheet, fila: int, texto: str, hasta: int, fuente: Font = VALOR) -> None:
    """Texto largo en una fila combinada (Header, Ejemplo), con alto según las líneas."""
    for c in range(1, hasta + 1):
        hoja.cell(fila, c).border = BORDE
    _celda(hoja, fila, 1, texto, fuente)
    if hasta > 1:
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=hasta)
    lineas = texto.count("\n") + 1
    hoja.row_dimensions[fila].height = min(409, max(15, 14 * lineas))


def _nombre_hoja(base: str, usados: set[str]) -> str:
    limpio = _PROHIBIDOS.sub(" ", base).strip().strip("'") or "Servicio"
    candidato, n = limpio[:31], 2
    while candidato.lower() in usados:
        sufijo = f" ({n})"
        candidato, n = limpio[: 31 - len(sufijo)] + sufijo, n + 1
    usados.add(candidato.lower())
    return candidato


def _marca(hoja: Worksheet, documento: Documento) -> None:
    hoja.oddHeader.right.text = _limpio(documento.marca_agua)
    hoja.oddHeader.right.size = 8
    hoja.oddFooter.left.text = f"SHA-256 {documento.hash}"
    hoja.oddFooter.left.size = 7
    hoja.oddFooter.right.text = "Página &P de &N"
    hoja.oddFooter.right.size = 8


def _tabla_derecha(hoja: Worksheet, fila: int, col: int, titulo: str, columnas: Sequence[str], filas) -> int:
    """Una tabla de referencia a la derecha de la hoja; devuelve la fila siguiente libre."""
    ancho = max(1, len(columnas))
    _encabezado(hoja, fila, col, titulo, col + ancho - 1)
    for j, nombre in enumerate(columnas):
        _encabezado(hoja, fila + 1, col + j, nombre)
    for i, valores in enumerate(filas, start=fila + 2):
        for j, valor in enumerate(valores):
            _celda(hoja, i, col + j, valor)
    return fila + 2 + len(filas) + 1


def _tabla_referencia(hoja: Worksheet, fila: int, col: int, tabla: Tabla) -> int:
    return _tabla_derecha(hoja, fila, col, tabla.titulo, tabla.columnas, tabla.filas)


def _matriz_documentos(hoja: Worksheet, fila: int, col: int, m: Matriz) -> int:
    columnas = ["ID", "NOMBRE", *(t.nombre for t in m.tipos)]
    valores = {"obligatorio": "S", "opcional": "N", "no_aplica": ""}
    filas = [[f.campo, f.detalle, *(valores[f.valores[t.clave]] for t in m.tipos)] for f in m.filas]
    return _tabla_derecha(hoja, fila, col, "Tabla guía documentos adjuntos", columnas, filas)


def _observacion(c: Campo) -> str:
    donde = {"ruta": "En la ruta", "consulta": "En la consulta", "cabecera": "En la cabecera", "cookie": "Cookie"}
    partes = [donde[c.ubicacion]] if c.ubicacion in donde else []
    if c.restricciones:
        partes.append(c.restricciones)
    return "; ".join(partes)


def _hoja_servicio(hoja: Worksheet, contrato: Contrato, op: Operacion, documento: Documento) -> None:
    m = op.matriz
    columnas = (
        ["Campo", "Descripción", *(f"OBLIGATORIO {t.nombre.upper()}" for t in m.tipos), "Tipo de Dato", "Observaciones"]
        if m
        else ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad", "Observaciones"]
    )
    ancho = len(columnas)
    col_meta = ancho - 1  # «Versión», «Fecha» y «Contrato» a la derecha, como en el original

    # Ambiente URL, versión, fecha y contrato
    _encabezado(hoja, 1, 1, "Ambiente URL", 2)
    ambientes = contrato.ambientes or []
    for i, a in enumerate(ambientes, start=2):
        _celda(hoja, i, 1, f"Ambiente {a.nombre}", ETIQUETA)
        _celda(hoja, i, 2, a.url)
    for i, (etiqueta, valor) in enumerate(
        (("Versión", contrato.version), ("Fecha", documento.fecha), ("Contrato", contrato.titulo)), start=1
    ):
        _encabezado(hoja, i, col_meta, etiqueta)
        _celda(hoja, i, col_meta + 1, valor)
    fila = max(2 + len(ambientes), 4) + 1

    _encabezado(hoja, fila, 1, "Header", 2)
    _bloque(hoja, fila + 1, "\n".join(op.cabeceras), 2)
    fila += 2
    _encabezado(hoja, fila, 1, "Método", 2)
    _celda(hoja, fila + 1, 1, op.metodo.capitalize(), CAMPO)
    _celda(hoja, fila + 1, 2, op.ruta)
    fila += 2
    if op.descripcion or op.resumen:
        _encabezado(hoja, fila, 1, "Descripción", ancho)
        _bloque(hoja, fila + 1, op.descripcion or op.resumen, ancho)
        fila += 2
    fila += 1

    # Parámetros
    _encabezado(hoja, fila, 1, "Parámetros", ancho)
    for j, nombre in enumerate(columnas, start=1):
        _encabezado(hoja, fila + 1, j, nombre)
    fila += 2
    campos = [*op.parametros, *op.cuerpo]
    if not campos:
        _celda(hoja, fila, 1, "Sin parámetros")
        fila += 1
    for c in campos:
        f = m.fila(c.nombre) if m else None
        _celda(hoja, fila, 1, c.nombre, CAMPO)
        _celda(hoja, fila, 2, c.descripcion)
        if m:
            for j, t in enumerate(m.tipos, start=3):
                sigla = SIGLAS_OBLIGATORIEDAD[f.valores[t.clave]] if f else ("S" if c.obligatorio else "N")
                _celda(hoja, fila, j, sigla)
            _celda(hoja, fila, 3 + len(m.tipos), c.tipo)
            observacion = "; ".join(x for x in ((f.condicion if f else ""), _observacion(c)) if x)
            _celda(hoja, fila, 4 + len(m.tipos), observacion)
        else:
            _celda(hoja, fila, 3, c.tipo)
            _celda(hoja, fila, 4, "S" if c.obligatorio else "N")
            _celda(hoja, fila, 5, _observacion(c))
        fila += 1
    if op.ejemplo_entrada:
        _encabezado(hoja, fila, 1, "Ejemplo", ancho)
        _bloque(hoja, fila + 1, op.ejemplo_entrada, ancho, CODIGO)
        fila += 2
    fila += 1

    # Salida
    for r in op.respuestas:
        if not (r.campos or r.ejemplo):
            continue
        _encabezado(hoja, fila, 1, f"Salida {r.codigo}", ancho)
        fila += 1
        if r.campos:
            for j, nombre in enumerate(("Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"), start=1):
                _encabezado(hoja, fila, j, nombre)
            fila += 1
            for c in r.campos:
                _celda(hoja, fila, 1, c.nombre, CAMPO)
                _celda(hoja, fila, 2, c.descripcion)
                _celda(hoja, fila, 3, c.tipo)
                _celda(hoja, fila, 4, "S" if c.obligatorio else "N")
                fila += 1
        if r.ejemplo:
            _encabezado(hoja, fila, 1, "Ejemplo", ancho)
            _bloque(hoja, fila + 1, r.ejemplo, ancho, CODIGO)
            fila += 2
        fila += 1
    if op.respuestas:
        _encabezado(hoja, fila, 1, "Código", 1)
        _encabezado(hoja, fila, 2, "Descripción", ancho)
        fila += 1
        for r in op.respuestas:
            _celda(hoja, fila, 1, r.codigo, CAMPO)
            _celda(hoja, fila, 2, r.descripcion)
            if ancho > 2:
                hoja.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=ancho)
            fila += 1

    # Tablas de referencia a la derecha
    col_tablas = ancho + 3
    fila_tabla = 6
    if op.documentos:
        fila_tabla = _matriz_documentos(hoja, fila_tabla, col_tablas, op.documentos)
    for t in op.tablas:
        fila_tabla = _tabla_referencia(hoja, fila_tabla, col_tablas, t)

    anchos = [34, 60, *([16] * len(m.tipos)), 18, 40] if m else [34, 60, 18, 16, 40]
    for j, w in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(j)].width = w
    for j in range(col_tablas, col_tablas + 8):
        hoja.column_dimensions[get_column_letter(j)].width = 22
    _marca(hoja, documento)


def _portada(hoja: Worksheet, documento: Documento) -> None:
    _celda(hoja, 1, 1, "FLIT · ATLAS", Font(name="Arial", size=12, bold=True, color="FF476BD9"), borde=False)
    _celda(hoja, 3, 1, documento.paquete, Font(name="Arial", size=18, bold=True, color="FF162744"), borde=False)
    _celda(hoja, 4, 1, f"Paquete de contratos de integración · revisión {documento.revision}", borde=False)
    filas = [
        ("Organización", documento.organizacion),
        ("Fecha", f"{documento.fecha} (hora de Bogotá)"),
        ("Revisión", str(documento.revision)),
        ("Contratos", "; ".join(f"{c.titulo} {c.version} ({c.tipo})" for c in documento.contratos)),
        ("SHA-256", documento.hash),
        ("Marca de agua", documento.marca_agua),
    ]
    for i, (etiqueta, valor) in enumerate(filas, start=6):
        _encabezado(hoja, i, 1, etiqueta)
        _celda(hoja, i, 2, valor)
    _celda(hoja, 13, 1, documento.leyenda, Font(name="Calibri", size=11, italic=True), borde=False)
    hoja.column_dimensions["A"].width = 22
    hoja.column_dimensions["B"].width = 90


def _indice(hoja: Worksheet, documento: Documento, servicios: list[tuple[str, Operacion]]) -> None:
    _encabezado(hoja, 1, 1, f"Versión {documento.revision}", 2)
    _encabezado(hoja, 2, 1, f"Fecha actualización: {documento.fecha}", 2)
    _encabezado(hoja, 3, 1, "Listado de servicios integrados", 2)
    for i, (nombre_hoja, op) in enumerate(servicios, start=4):
        _celda(hoja, i, 1, nombre_hoja, CAMPO)
        celda = hoja.cell(i, 1)
        celda.hyperlink = f"#'{nombre_hoja}'!A1"
        celda.font = Font(name="Calibri", size=11, bold=True, color="FF476BD9", underline="single")
        _celda(hoja, i, 2, f"{op.titulo} ({op.metodo} {op.ruta})")
    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 90


def _cambios(hoja: Worksheet, documento: Documento) -> None:
    for j, nombre in enumerate(("Revisión", "Fecha", "Cambio"), start=1):
        _encabezado(hoja, 1, j, nombre)
    fila = 2
    for e in documento.changelog or []:
        _celda(hoja, fila, 1, e.revision)
        _celda(hoja, fila, 2, e.fecha)
        _celda(hoja, fila, 3, e.cambio)
        fila += 1
    if not documento.changelog:
        _celda(hoja, fila, 1, "Primera revisión del paquete.")
        fila += 1
    for c in documento.contratos:
        if not c.historial:
            continue
        fila += 1
        _encabezado(hoja, fila, 1, f"Historial de {c.titulo}", 3)
        fila += 1
        for j, nombre in enumerate(("Versión", "Fecha", "Cambio"), start=1):
            _encabezado(hoja, fila, j, nombre)
        fila += 1
        for h in c.historial:
            _celda(hoja, fila, 1, h["version"])
            _celda(hoja, fila, 2, h["fecha"])
            _celda(hoja, fila, 3, h["cambio"])
            fila += 1
    hoja.column_dimensions["A"].width = 16
    hoja.column_dimensions["B"].width = 14
    hoja.column_dimensions["C"].width = 100


def _tablas(hoja: Worksheet, documento: Documento) -> None:
    fila = 1
    for c in documento.contratos:
        for m in c.matrices:
            if m.clase == "documentos":
                fila = _matriz_documentos(hoja, fila, 1, m)
            else:
                columnas = ["Campo", *(t.nombre for t in m.tipos), "Observaciones"]
                filas = [
                    [f.campo, *(SIGLAS_OBLIGATORIEDAD[f.valores[t.clave]] for t in m.tipos), f.condicion]
                    for f in m.filas
                ]
                fila = _tabla_derecha(hoja, fila, 1, f"{m.titulo} ({m.propiedad})", columnas, filas)
        for t in c.tablas:
            fila = _tabla_referencia(hoja, fila, 1, t)
    if fila == 1:
        _celda(hoja, 1, 1, "Los contratos del paquete no traen tablas de referencia.")
    hoja.column_dimensions["A"].width = 48
    for j in range(2, 10):
        hoja.column_dimensions[get_column_letter(j)].width = 20


def generar(documento: Documento) -> bytes:
    libro = Workbook()
    portada = libro.worksheets[0]
    portada.title = "Portada"
    indice = libro.create_sheet("Índice")
    cambios = libro.create_sheet("Cambios")
    usados = {h.lower() for h in HOJAS_FIJAS}
    servicios: list[tuple[str, Operacion]] = []
    n = 0
    for contrato in documento.contratos:
        for op in contrato.operaciones:
            n += 1
            nombre = _nombre_hoja(f"{n} - {op.operation_id}", usados)
            _hoja_servicio(libro.create_sheet(nombre), contrato, op, documento)
            servicios.append((nombre, op))
    _tablas(libro.create_sheet("Tablas de referencia"), documento)
    if documento.guias:
        guias = libro.create_sheet("Guías")
        _encabezado(guias, 1, 1, "Guía")
        _encabezado(guias, 1, 2, "Contenido (Markdown)")
        for i, g in enumerate(documento.guias, start=2):
            _celda(guias, i, 1, g.titulo, CAMPO)
            _celda(guias, i, 2, g.markdown)
        guias.column_dimensions["A"].width = 34
        guias.column_dimensions["B"].width = 120
    _portada(portada, documento)
    _indice(indice, documento, servicios)
    _cambios(cambios, documento)
    for hoja in (portada, indice, cambios):
        _marca(hoja, documento)

    # Propiedades fijas (sin la hora de generación) y guardado sin `save_workbook`, que sella la hora actual.
    fecha = datetime.combine(documento.dia, time())
    libro.properties.creator = "Flit Atlas"
    libro.properties.lastModifiedBy = "Flit Atlas"
    libro.properties.title = _limpio(documento.titulo)
    libro.properties.subject = _limpio(documento.marca_agua)
    libro.properties.keywords = documento.hash
    libro.properties.created = fecha
    libro.properties.modified = fecha
    salida = io.BytesIO()
    with ZipFile(salida, "w", ZIP_DEFLATED, allowZip64=True) as archivo:
        ExcelWriter(libro, archivo).save()
    return normalizar_zip(salida.getvalue(), documento.dia)
