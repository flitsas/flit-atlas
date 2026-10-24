"""Lectura del Excel de contratos ICT (todas las hojas) a un modelo neutro.

Cada hoja de contrato tiene la misma forma: URL por ambiente, versión, fecha y colaborador, cabeceras, método, tabla de
parámetros (con columnas de obligatoriedad por tipo de trámite en la de transacciones), ejemplo de entrada, tabla de
salida, ejemplos de salida, tabla de códigos de respuesta y, a la derecha, tablas de referencia («Tabla …»). La
lectura no conoce los nombres de las hojas: se guía por esas etiquetas. Lo que no reconoce queda en
`celdas_sin_interpretar` para que el informe lo muestre.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

AMBIENTES = {"ambiente dev": "dev", "ambiente qa": "qa", "ambiente pdn": "pdn"}
_CODIGO = re.compile(r"^(\d{3})\s+(.*)$")
_ETIQUETAS = {"header", "metodo", "ejemplo", "salida", "campo", "ambiente url"}


@dataclass
class Campo:
    hoja: str
    celda: str
    seccion: str  # "parametros" | "salida"
    nombre: str
    descripcion: str | None = None
    tipo: str | None = None
    obligatorio: str | None = None  # "S" | "N" | None
    por_tipo: dict[str, str] = field(default_factory=dict)  # columna «OBLIGATORIO …» → "S" | "N"
    observacion: str | None = None

    @property
    def ruta(self) -> list[str]:
        return [parte.strip() for parte in self.nombre.split("/") if parte.strip()]


@dataclass
class Tabla:
    hoja: str
    celda: str
    titulo: str
    columnas: list[str]
    filas: list[list[str | None]]


@dataclass
class Codigo:
    hoja: str
    celda: str
    codigo: int
    texto: str
    descripcion: str | None


@dataclass
class Hoja:
    nombre: str
    version: str | None = None
    fecha: date | None = None
    colaborador: str | None = None
    urls: dict[str, str | None] = field(default_factory=dict)
    cabeceras: str | None = None
    metodo: str | None = None
    listado: bool = False
    campos: list[Campo] = field(default_factory=list)
    ejemplo_entrada: str | None = None
    ejemplos_salida: list[str] = field(default_factory=list)
    codigos: list[Codigo] = field(default_factory=list)
    tablas: list[Tabla] = field(default_factory=list)
    leyenda: dict[str, str] = field(default_factory=dict)  # «1» → «Obligatorio» (guía de adjuntos)
    celdas_sin_interpretar: list[tuple[str, str]] = field(default_factory=list)

    @property
    def servicio(self) -> str:
        """«2 - Transacction» → «Transacction»."""
        return re.sub(r"^\s*\d+\s*-\s*", "", self.nombre).strip()


@dataclass
class Indice:
    hoja: str
    version: str | None
    fecha: str | None
    titulo: str | None
    servicios: list[tuple[str, str | None]]


@dataclass
class Libro:
    archivo: str  # solo el nombre, nunca la ruta local
    indice: Indice | None
    hojas: list[Hoja]
    hojas_vacias: list[str]

    def hoja(self, servicio: str) -> Hoja | None:
        return next((h for h in self.hojas if h.servicio.lower() == servicio.lower()), None)


def _texto(valor: object) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    return texto or None


def _celda(fila: int, columna: int) -> str:
    return f"{get_column_letter(columna)}{fila}"


class _Rejilla:
    def __init__(self, hoja) -> None:
        self.valores: dict[tuple[int, int], object] = {}
        for fila in hoja.iter_rows():
            for c in fila:
                if c.value is not None and _texto(c.value) is not None:
                    self.valores[(c.row, c.column)] = c.value
        self.usadas: set[tuple[int, int]] = set()
        self.max_fila = max((f for f, _ in self.valores), default=0)
        self.max_col = max((c for _, c in self.valores), default=0)

    def texto(self, fila: int, columna: int) -> str | None:
        return _texto(self.valores.get((fila, columna)))

    def usar(self, fila: int, columna: int) -> str | None:
        if (fila, columna) in self.valores:
            self.usadas.add((fila, columna))
        return self.texto(fila, columna)


def _leer_tablas(rejilla: _Rejilla, nombre_hoja: str) -> tuple[list[Tabla], int]:
    """Tablas de referencia: un título «Tabla …», una fila de encabezado y filas hasta el siguiente título."""
    titulos = sorted(
        (f, c)
        for (f, c), v in rejilla.valores.items()
        if isinstance(v, str) and v.strip().lower().startswith("tabla ") and c > 1
    )
    tablas: list[Tabla] = []
    for f, c in titulos:
        if (f, c) in rejilla.usadas:
            continue
        titulo = rejilla.usar(f, c) or ""
        encabezado = f + 1
        # Títulos repetidos en filas seguidas (p. ej. «Tabla Tipo Documento» dos veces).
        while (rejilla.texto(encabezado, c) or "").lower() == titulo.lower():
            rejilla.usar(encabezado, c)
            encabezado += 1
        columnas: list[str] = []
        col = c
        while rejilla.texto(encabezado, col) is not None:
            columnas.append(rejilla.usar(encabezado, col) or "")
            col += 1
        siguiente = min((ft for ft, ct in titulos if ct == c and ft > encabezado), default=rejilla.max_fila + 1)
        filas = []
        for fila in range(encabezado + 1, siguiente):
            valores = [rejilla.texto(fila, c + i) for i in range(len(columnas))]
            if any(v is not None for v in valores):
                for i in range(len(columnas)):
                    rejilla.usar(fila, c + i)
                filas.append(valores)
        tablas.append(Tabla(nombre_hoja, _celda(f, c), titulo, columnas, filas))
    limite = min((c for _, c in titulos), default=rejilla.max_col + 1)
    return tablas, limite


def _leer_contrato(rejilla: _Rejilla, nombre: str) -> Hoja:
    hoja = Hoja(nombre)
    for (f, c), v in list(rejilla.valores.items()):
        etiqueta = (_texto(v) or "").lower()
        if etiqueta in ("version", "fecha", "colaborador"):
            rejilla.usar(f, c)
            valor = rejilla.valores.get((f, c + 1))
            rejilla.usar(f, c + 1)
            if etiqueta == "version":
                hoja.version = _texto(valor)
            elif etiqueta == "fecha":
                hoja.fecha = valor.date() if isinstance(valor, datetime) else None
            else:
                hoja.colaborador = _texto(valor)
    hoja.tablas, limite = _leer_tablas(rejilla, nombre)

    seccion = None
    fila = 1
    while fila <= rejilla.max_fila:
        etiqueta = rejilla.texto(fila, 1)
        clave = (etiqueta or "").lower()
        if etiqueta is None:
            fila += 1
            continue
        if clave in AMBIENTES:
            rejilla.usar(fila, 1)
            hoja.urls[AMBIENTES[clave]] = rejilla.usar(fila, 2)
        elif clave == "ambiente url":
            rejilla.usar(fila, 1)
        elif clave in ("header", "metodo"):
            rejilla.usar(fila, 1)
            siguiente = fila + 1
            valor = rejilla.usar(siguiente, 1)
            if clave == "header":
                hoja.cabeceras = valor
            else:
                hoja.metodo = (valor or "").lower() or None
            fila = siguiente
        elif clave.startswith("parametros"):
            rejilla.usar(fila, 1)
            seccion = "parametros"
            hoja.listado = "listado" in clave
        elif clave == "salida":
            rejilla.usar(fila, 1)
            seccion = "salida"
        elif clave == "ejemplo":
            rejilla.usar(fila, 1)
            siguiente = fila + 1
            while siguiente <= rejilla.max_fila and rejilla.texto(siguiente, 1) is None:
                siguiente += 1
            valor = rejilla.usar(siguiente, 1)
            if seccion == "parametros":
                hoja.ejemplo_entrada = valor
            else:
                hoja.ejemplos_salida.append(valor or "")
            fila = siguiente
        elif clave == "campo":
            fila = _leer_tabla_campos(rejilla, hoja, fila, limite, seccion or "parametros")
            continue
        elif seccion == "salida" and etiqueta[:1] in "[{":
            # Ejemplos de error sueltos después de un «Salida» sin tabla.
            rejilla.usar(fila, 1)
            hoja.ejemplos_salida.append(etiqueta)
        fila += 1

    # Leyenda suelta junto a una tabla («1 Obligatorio», «2 Opcional»).
    for (f, c), v in list(rejilla.valores.items()):
        if (f, c) not in rejilla.usadas and (_texto(v) or "").lower() in ("obligatorio", "opcional"):
            clave = rejilla.texto(f, c - 1)
            if clave and clave.isdigit() and (f, c - 1) not in rejilla.usadas:
                hoja.leyenda[clave] = _texto(v) or ""
                rejilla.usar(f, c)
                rejilla.usar(f, c - 1)

    hoja.celdas_sin_interpretar = sorted(
        ((_celda(f, c), _texto(v) or "") for (f, c), v in rejilla.valores.items() if (f, c) not in rejilla.usadas),
        key=lambda x: (int(re.sub(r"\D", "", x[0])), x[0]),
    )
    return hoja


def _leer_tabla_campos(rejilla: _Rejilla, hoja: Hoja, fila: int, limite: int, seccion: str) -> int:
    columnas: dict[int, str] = {}
    for col in range(1, limite):
        titulo = rejilla.usar(fila, col)
        if titulo:
            columnas[col] = titulo
    fila += 1
    while fila <= rejilla.max_fila:
        nombre = rejilla.texto(fila, 1)
        if nombre is None or nombre.lower() in _ETIQUETAS or nombre.lower().startswith("parametros"):
            break
        valores = {titulo: rejilla.usar(fila, col) for col, titulo in columnas.items()}
        rejilla.usar(fila, 1)
        if m := _CODIGO.match(nombre):
            descripcion = next((v for t, v in valores.items() if t.lower().startswith("descrip")), None)
            hoja.codigos.append(Codigo(hoja.nombre, _celda(fila, 1), int(m.group(1)), m.group(2), descripcion))
        else:
            campo = Campo(hoja.nombre, _celda(fila, 1), seccion, nombre)
            for titulo, valor in valores.items():
                t = titulo.lower()
                if t == "campo":
                    continue
                if t.startswith("descrip"):
                    campo.descripcion = valor
                elif t.startswith("tipo de dato"):
                    campo.tipo = valor
                elif t.startswith("obligatoriedad"):
                    campo.obligatorio = (valor or "").upper() or None
                elif t.startswith("obligatorio "):
                    campo.por_tipo[titulo[len("obligatorio ") :].strip()] = (valor or "").upper()
                elif t.startswith("observ"):
                    campo.observacion = valor
            hoja.campos.append(campo)
        fila += 1
    return fila


def _leer_indice(rejilla: _Rejilla, nombre: str) -> Indice:
    version = fecha = titulo = None
    servicios = []
    for fila in range(1, rejilla.max_fila + 1):
        a, b = rejilla.usar(fila, 1), rejilla.usar(fila, 2)
        if a is None:
            continue
        if a.lower().startswith("version"):
            version = a.split(" ", 1)[1].strip() if " " in a else None
        elif a.lower().startswith("fecha"):
            fecha = a.split(":", 1)[1].strip() if ":" in a else None
        elif b is None and titulo is None and len(a) > 30:
            titulo = a
        else:
            servicios.append((a, b))
    return Indice(nombre, version, fecha, titulo, servicios)


def leer_libro(ruta: Path) -> Libro:
    libro = load_workbook(ruta, data_only=True)
    indice, hojas, vacias = None, [], []
    for hoja in libro.worksheets:
        rejilla = _Rejilla(hoja)
        if not rejilla.valores:
            vacias.append(hoja.title)
        elif any((_texto(v) or "").lower() == "ambiente url" for v in rejilla.valores.values()):
            hojas.append(_leer_contrato(rejilla, hoja.title))
        elif indice is None:
            indice = _leer_indice(rejilla, hoja.title)
    return Libro(Path(ruta).name, indice, hojas, vacias)
