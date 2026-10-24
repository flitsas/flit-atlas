"""Fixtures ficticias del enriquecimiento de core-ict: un repo `flit` mínimo (contrato y Gateway) y un xlsx con la
forma del Excel.

Nada de aquí sale del Excel real: los valores sensibles de prueba (token con forma de JWT, correo de un dominio
ajeno, ruta local, IP privada) se arman en tiempo de ejecución para que no queden literales en git.
"""

import base64
import json
from datetime import datetime
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook


def _b64(texto: str) -> str:
    return base64.urlsafe_b64encode(texto.encode()).decode().rstrip("=")


# Valores sensibles ficticios, armados en tiempo de ejecución.
TOKEN_FICTICIO = ".".join([_b64('{"alg":"HS256","typ":"JWT"}'), _b64('{"sub":"prueba"}'), "firmaDePrueba123"])
CORREO_AJENO = "cliente" + "@" + "empresa-ficticia.test"
RUTA_LOCAL = "C:" + "\\Users\\persona\\Descargas\\soporte.pdf"
IP_PRIVADA = ".".join(["10", "20", "30", "40"])
COLABORADOR = "Persona Ficticia"

SENSIBLES = [TOKEN_FICTICIO, CORREO_AJENO, RUTA_LOCAL, IP_PRIVADA, COLABORADOR]


def _metadatos(hoja, col_etiqueta: int, version: str, url_qa: str | None) -> None:
    hoja.cell(1, 1, "Ambiente URL")
    hoja.cell(2, 1, "Ambiente DEV")
    hoja.cell(3, 1, "Ambiente QA")
    if url_qa:
        hoja.cell(3, 2, url_qa)
    hoja.cell(4, 1, "Ambiente PDN")
    hoja.cell(1, col_etiqueta, "Version")
    hoja.cell(1, col_etiqueta + 1, version)
    hoja.cell(2, col_etiqueta, "Fecha")
    hoja.cell(2, col_etiqueta + 1, datetime(2025, 5, 16))
    hoja.cell(3, col_etiqueta, "Colaborador")
    hoja.cell(3, col_etiqueta + 1, COLABORADOR)


def _filas(hoja, inicio: int, filas: list[list]) -> int:
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            if valor is not None:
                hoja.cell(inicio + i, j + 1, valor)
    return inicio + len(filas)


def crear_excel(ruta: Path) -> Path:
    libro = Workbook()
    indice = libro.active
    indice.title = "Hoja1"
    _filas(
        indice,
        1,
        [
            ["Version 1.0"],
            ["Fecha actualizacion: 01-01-2025"],
            ["Listado de servicios integracion terceros gestores"],
            ["Auth", "Servicio de autenticacion de prueba"],
            ["Transacction", "Servicio de registro de tramites de prueba"],
            ["WebHookTransaction"],
            ["RequestSign", "Servicio sin hoja"],
        ],
    )
    libro.create_sheet("Hoja2")

    auth = libro.create_sheet("1 - Auth")
    _metadatos(auth, 4, "1.2", "https://ict.qa.flitsas.online/api/v1/auth/login")
    _filas(
        auth,
        5,
        [
            ["Header"],
            ["Accept: application/json\nContent-Type: application/json"],
            ["Metodo"],
            ["Post"],
            [],
            ["Parametros"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["Usuario", "Correo del gestor", "String(100)", "S"],
            ["Password", "Clave del gestor", "String", "S"],
            ["Ejemplo"],
            ['{\n  "username": "' + CORREO_AJENO + '",\n  "password ": "clave-de-prueba"\n}'],
            [],
            ["Salida"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["Token", "Token para las demás operaciones", "String"],
            ["Ejemplo"],
            ['{\n  "token": “' + TOKEN_FICTICIO + "”\n}"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["401 Unauthorized", "No autorizado.", "String"],
            ["404 Not Found", "No encontrado.", "String"],
            ["500 Internal Server Error", "Error interno.", "String"],
            ["Salida"],
            ['{"message": "Incorrect username or password."}\n\n{"message": "An unexpected error occurred."}'],
        ],
    )

    tramite = libro.create_sheet("2 - Transacction")
    _metadatos(tramite, 8, "3", "https://ict.qa.flitsas.online/api/v1/external-transaction/register")
    cabecera = [
        "Campo",
        "Descripción",
        "OBLIGATORIO TRASPASO BILATERAL",
        "OBLIGATORIO MATRICULA INICIAL",
        "Tipo de Dato",
        "OBSERVACIONES",
    ]
    _filas(
        tramite,
        5,
        [
            ["Header"],
            ["Accept: application/json\nAuthorization: Bearer access_token"],
            ["Metodo"],
            ["Post"],
            [],
            ["Parametros [Listado]"],
            cabecera,
            ["company_manager_id", "Id de la compañía gestora", "S", "S", "number"],
            ["transaction_type", "Tipo de Tramite (Ver Tabla Tipo Tramite)", "S", "S", "number(2)"],
            ["plate", "Placa del vehículo", "S", "N", "String(6)"],
            ["vin", "VIN del vehículo", "N", "S", "String(17)", "Es obligatorio si es Matricula"],
            ["seller/document_type", "Tipo de documento del vendedor", "S", "S", "String(5)"],
            ["seller/email", "Correo del vendedor", "S", "S", "String(255)"],
            ["buyer/name", "Nombre del comprador", "S", "N", "String(100)"],
            ["UrlWebHook", "URL de notificación", "N", "N", "String(100)"],
            ["Ejemplo"],
            [
                json.dumps(
                    [
                        {
                            "company_manager_id": 7,
                            "transaction_type": 3,
                            "plate": "QWE987",
                            "seller": [{"document_type": "CC", "email": CORREO_AJENO}],
                            "buyer": [{"name": "Persona Real"}],
                            "url_web_hook": "http://" + IP_PRIVADA + "/notificar",
                        }
                    ],
                    indent=2,
                ).replace("}\n  ]", "},\n  ]")
            ],
            [],
            ["Salida"],
            ["Campo", "Descripción", None, None, "Tipo de Dato", "Obligatoriedad"],
            ["TotalRows", "Registros recibidos", None, None, "Integer"],
            ["Detail/Plate", "Placa", None, None, "String(15)"],
            ["Ejemplo"],
            ['{"TotalRows": 1, "Detail": [{"Plate": "QWE987", "Status": 1, "Message": "ok", "TransactionFlit": 5}]}'],
        ],
    )
    _filas(tramite, 6, [[None] * 12 + ["Tabla Tipo Documento"]])
    _filas(tramite, 7, [[None] * 12 + ["ID", "document_type_code", "NOMBRE"]])
    _filas(tramite, 8, [[None] * 12 + ["1", "NIT", "NIT"], [None] * 12 + ["2", "CC", "Cedula"]])
    _filas(tramite, 12, [[None] * 12 + ["Tabla Tipo Tramite"], [None] * 12 + ["ID", "NOMBRE"]])
    _filas(
        tramite,
        14,
        [
            [None] * 12 + ["1", "Matricula Inicial"],
            [None] * 12 + ["3", "Traspaso Bilateral"],
            [None] * 12 + ["5", "Blindaje"],
        ],
    )

    webhook = libro.create_sheet("4 - WebHookTransaction")
    for fila, texto in enumerate(["Ambiente URL", "Ambiente DEV", "Ambiente QA", "Ambiente PDN"], start=1):
        webhook.cell(fila, 1, texto)
    _filas(
        webhook,
        5,
        [
            ["Header"],
            ["Authorization: Bearer access_token"],
            ["Metodo"],
            ["Post"],
            [],
            ["Parametros"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["TransactionFlit", "Id del trámite", "number", "S"],
            ["statusMessage", "Mensaje del estado", "String(250)", "N"],
            ["Ejemplo"],
            ['{"transactionFlit": 10, "statusMessage": "Ver ' + RUTA_LOCAL.replace("\\", "\\\\") + '",}'],
            ["Salida"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["Status", "Estado de recepción", "Integer"],
            ["Ejemplo"],
            ['{"Status": 1}'],
        ],
    )
    libro.save(ruta)
    return ruta


def _contrato_core_ict() -> dict:
    """Un core-ict.v1.yaml mínimo y ficticio con la forma del de FLIT 2 (OpenAPI 3.0.3, sin operationId)."""
    error = {"description": "Token ICT ausente o inválido."}
    return {
        "openapi": "3.0.3",
        "info": {
            "title": "FLIT Core ICT API",
            "version": "1.0.0",
            "description": "Contrato de prueba de core-ict. Notas en " + RUTA_LOCAL,
        },
        "servers": [
            {"url": "https://api.ejemplo-viejo.test", "description": "Producción"},
            {"url": "http://" + "local" + "host:4002", "description": "Local"},
        ],
        "tags": [
            {"name": "ICT · Auth", "description": "Login del cliente de integración."},
            {"name": "ICT · Ingesta", "description": "Registro por lote."},
            {"name": "ICT · Estado", "description": "Consulta de estado."},
            {"name": "ICT · Observabilidad", "description": "Logs internos."},
        ],
        "paths": {
            "/api/v1/auth/login": {
                "post": {
                    "tags": ["ICT · Auth"],
                    "summary": "Login",
                    "security": [],
                    "requestBody": {
                        "required": True,
                        "content": {"application/json": {"schema": {"$ref": "#/components/schemas/IctLoginRequest"}}},
                    },
                    "responses": {
                        "200": {
                            "description": "Token emitido.",
                            "content": {
                                "application/json": {"schema": {"$ref": "#/components/schemas/IctLoginResponse"}}
                            },
                        },
                        "401": {"description": "Credenciales inválidas."},
                    },
                }
            },
            "/api/v1/external-transaction/register": {
                "post": {
                    "tags": ["ICT · Ingesta"],
                    "summary": "Registro por lote",
                    "security": [{"ictClientToken": []}],
                    "requestBody": {
                        "required": True,
                        "content": {
                            "application/json": {"schema": {"$ref": "#/components/schemas/RegisterBatchRequest"}}
                        },
                    },
                    "responses": {
                        "200": {
                            "description": "Lote procesado.",
                            "content": {
                                "application/json": {"schema": {"$ref": "#/components/schemas/RegisterBatchResponse"}}
                            },
                        },
                        "401": error,
                    },
                }
            },
            "/api/v1/status-process/byId/{id}": {
                "get": {
                    "tags": ["ICT · Estado"],
                    "summary": "Estado v1",
                    "security": [{"ictClientToken": []}],
                    "parameters": [{"$ref": "#/components/parameters/ManagerIdTransaction"}],
                    "responses": {
                        "200": {
                            "description": "Estado.",
                            "content": {
                                "application/json": {"schema": {"$ref": "#/components/schemas/StatusProcessV1Response"}}
                            },
                        },
                        "404": {"description": "No existe."},
                    },
                }
            },
            "/api/v1/ict/logs": {
                "get": {
                    "tags": ["ICT · Observabilidad"],
                    "summary": "Logs",
                    "security": [{"platformJwt": []}],
                    "responses": {"200": {"description": "Página de logs."}},
                }
            },
        },
        "components": {
            "securitySchemes": {
                "ictClientToken": {"type": "http", "scheme": "bearer", "bearerFormat": "JWT"},
                "platformJwt": {"type": "http", "scheme": "bearer", "bearerFormat": "JWT"},
            },
            "parameters": {
                "ManagerIdTransaction": {"name": "id", "in": "path", "required": True, "schema": {"type": "string"}}
            },
            "schemas": {
                "IctLoginRequest": {
                    "type": "object",
                    "required": ["username", "password"],
                    "properties": {"username": {"type": "string"}, "password": {"type": "string"}},
                },
                "IctLoginResponse": {
                    "type": "object",
                    "properties": {"token": {"type": "string"}, "expiresAt": {"type": "string"}},
                },
                "RegisterBatchRequest": {
                    "type": "object",
                    "properties": {"transactions": {"type": "array", "items": {"type": "object"}}},
                },
                "RegisterBatchResponse": {
                    "type": "object",
                    "properties": {
                        "TotalRows": {"type": "integer"},
                        "Detail": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "Plate": {"type": "string"},
                                    "Status": {"type": "integer"},
                                    "TransactionFlit": {"type": "string", "nullable": True},
                                },
                            },
                        },
                    },
                },
                "StatusProcessV1Response": {
                    "type": "object",
                    "properties": {
                        "statusValidation": {"type": "integer"},
                        "statusProcess": {"type": "array", "items": {"type": "object"}},
                    },
                },
            },
        },
    }


# appsettings del Gateway con comentarios, como el de FLIT 2.
GATEWAY = """{
  // Rutas del Gateway (ficticias)
  "ReverseProxy": {
    "Routes": {
      "ict-host-route": {
        "ClusterId": "core-ict-cluster",
        "Match": {
          "Hosts": ["ict.flitsas.com", "ict.dev.flitsas.com", "ict.qa.flitsas.com", "ict.HOST_LOCAL"],
          "Path": "/{**catch-all}"
        }
      },
      "otra-route": { "ClusterId": "core-api-cluster", "Match": { "Hosts": ["api.flitsas.com"] } },
    }
  }
}
""".replace("HOST_LOCAL", "local" + "host")


def crear_flit(carpeta: Path) -> Path:
    """Un repo `flit` ficticio con solo lo que lee la conversión: el contrato de core-ict y el Gateway."""
    contrato = carpeta / "contracts" / "openapi" / "core-ict.v1.yaml"
    contrato.parent.mkdir(parents=True)
    contrato.write_text(yaml.safe_dump(_contrato_core_ict(), allow_unicode=True, sort_keys=False), encoding="utf-8")
    gateway = carpeta / "services" / "core-api" / "src" / "Flit.Gateway" / "appsettings.json"
    gateway.parent.mkdir(parents=True)
    gateway.write_text(GATEWAY, encoding="utf-8")
    return carpeta


@pytest.fixture
def excel(tmp_path: Path) -> Path:
    return crear_excel(tmp_path / "Contratos_Integraciones_Prueba.xlsx")


@pytest.fixture
def flit(tmp_path: Path) -> Path:
    return crear_flit(tmp_path / "flit")
