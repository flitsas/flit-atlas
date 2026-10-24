"""Fixtures ficticias de la conversión del legado ICT: un swagger mínimo por repo y un xlsx con la forma del Excel.

Nada de aquí sale del Excel real: los valores sensibles de prueba (token con forma de JWT, correo de un dominio
ajeno, ruta local, IP privada) se arman en tiempo de ejecución para que no queden literales en git.
"""

import base64
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook


def _b64(texto: str) -> str:
    return base64.urlsafe_b64encode(texto.encode()).decode().rstrip("=")


# Valores sensibles ficticios, armados en tiempo de ejecución.
TOKEN_FICTICIO = ".".join([_b64('{"alg":"HS256","typ":"JWT"}'), _b64('{"sub":"prueba"}'), "firmaDePrueba123"])
CORREO_AJENO = "cliente" + "@" + "empresa-ficticia.test"
RUTA_LOCAL = "C:" + "\\Users\\persona\\Descargas\\soporte.pdf"
IP_PRIVADA = ".".join(["10", "20", "30", "40"])
COLABORADOR = "Persona Ficticia"

SENSIBLES = [TOKEN_FICTICIO, CORREO_AJENO, RUTA_LOCAL, IP_PRIVADA, COLABORADOR]


def _metadatos(hoja, col_etiqueta: int, version: str, url_qa: str | None) -> None:
    hoja.cell(1, 1, "Ambiente URL")
    hoja.cell(2, 1, "Ambiente DEV")
    hoja.cell(3, 1, "Ambiente QA")
    if url_qa:
        hoja.cell(3, 2, url_qa)
    hoja.cell(4, 1, "Ambiente PDN")
    hoja.cell(1, col_etiqueta, "Version")
    hoja.cell(1, col_etiqueta + 1, version)
    hoja.cell(2, col_etiqueta, "Fecha")
    hoja.cell(2, col_etiqueta + 1, datetime(2025, 5, 16))
    hoja.cell(3, col_etiqueta, "Colaborador")
    hoja.cell(3, col_etiqueta + 1, COLABORADOR)


def _filas(hoja, inicio: int, filas: list[list]) -> int:
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            if valor is not None:
                hoja.cell(inicio + i, j + 1, valor)
    return inicio + len(filas)


def crear_excel(ruta: Path) -> Path:
    libro = Workbook()
    indice = libro.active
    indice.title = "Hoja1"
    _filas(
        indice,
        1,
        [
            ["Version 1.0"],
            ["Fecha actualizacion: 01-01-2025"],
            ["Listado de servicios integracion terceros gestores"],
            ["Auth", "Servicio de autenticacion de prueba"],
            ["Transacction", "Servicio de registro de tramites de prueba"],
            ["WebHookTransaction"],
            ["RequestSign", "Servicio sin hoja"],
        ],
    )
    libro.create_sheet("Hoja2")

    auth = libro.create_sheet("1 - Auth")
    _metadatos(auth, 4, "1.2", "https://ict.qa.flitsas.online/api/v1/auth/login")
    _filas(
        auth,
        5,
        [
            ["Header"],
            ["Accept: application/json\nContent-Type: application/json"],
            ["Metodo"],
            ["Post"],
            [],
            ["Parametros"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["Usuario", "Correo del gestor", "String(100)", "S"],
            ["Password", "Clave del gestor", "String", "S"],
            ["Ejemplo"],
            ['{\n  "username": "' + CORREO_AJENO + '",\n  "password ": "clave-de-prueba"\n}'],
            [],
            ["Salida"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["Token", "Token para las demás operaciones", "String"],
            ["Ejemplo"],
            ['{\n  "token": “' + TOKEN_FICTICIO + "”\n}"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["401 Unauthorized", "No autorizado.", "String"],
            ["404 Not Found", "No encontrado.", "String"],
            ["500 Internal Server Error", "Error interno.", "String"],
            ["Salida"],
            ['{"message": "Incorrect username or password."}\n\n{"message": "An unexpected error occurred."}'],
        ],
    )

    tramite = libro.create_sheet("2 - Transacction")
    _metadatos(tramite, 8, "3", "https://ict.qa.flitsas.online/api/v1/external-transaction/register")
    cabecera = [
        "Campo",
        "Descripción",
        "OBLIGATORIO TRASPASO BILATERAL",
        "OBLIGATORIO MATRICULA INICIAL",
        "Tipo de Dato",
        "OBSERVACIONES",
    ]
    _filas(
        tramite,
        5,
        [
            ["Header"],
            ["Accept: application/json\nAuthorization: Bearer access_token"],
            ["Metodo"],
            ["Post"],
            [],
            ["Parametros [Listado]"],
            cabecera,
            ["company_manager_id", "Id de la compañía gestora", "S", "S", "number"],
            ["transaction_type", "Tipo de Tramite (Ver Tabla Tipo Tramite)", "S", "S", "number(2)"],
            ["plate", "Placa del vehículo", "S", "N", "String(6)"],
            ["vin", "VIN del vehículo", "N", "S", "String(17)", "Es obligatorio si es Matricula"],
            ["seller/document_type", "Tipo de documento del vendedor", "S", "S", "String(5)"],
            ["seller/email", "Correo del vendedor", "S", "S", "String(255)"],
            ["buyer/name", "Nombre del comprador", "S", "N", "String(100)"],
            ["UrlWebHook", "URL de notificación", "N", "N", "String(100)"],
            ["Ejemplo"],
            [
                json.dumps(
                    [
                        {
                            "company_manager_id": 7,
                            "transaction_type": 3,
                            "plate": "QWE987",
                            "seller": [{"document_type": "CC", "email": CORREO_AJENO}],
                            "buyer": [{"name": "Persona Real"}],
                            "url_web_hook": "http://" + IP_PRIVADA + "/notificar",
                        }
                    ],
                    indent=2,
                ).replace("}\n  ]", "},\n  ]")
            ],
            [],
            ["Salida"],
            ["Campo", "Descripción", None, None, "Tipo de Dato", "Obligatoriedad"],
            ["TotalRows", "Registros recibidos", None, None, "Integer"],
            ["Detail/Plate", "Placa", None, None, "String(15)"],
            ["Ejemplo"],
            ['{"TotalRows": 1, "Detail": [{"Plate": "QWE987", "Status": 1, "Message": "ok", "TransactionFlit": 5}]}'],
        ],
    )
    _filas(tramite, 6, [[None] * 12 + ["Tabla Tipo Documento"]])
    _filas(tramite, 7, [[None] * 12 + ["ID", "document_type_code", "NOMBRE"]])
    _filas(tramite, 8, [[None] * 12 + ["1", "NIT", "NIT"], [None] * 12 + ["2", "CC", "Cedula"]])
    _filas(tramite, 12, [[None] * 12 + ["Tabla Tipo Tramite"], [None] * 12 + ["ID", "NOMBRE"]])
    _filas(
        tramite,
        14,
        [
            [None] * 12 + ["1", "Matricula Inicial"],
            [None] * 12 + ["3", "Traspaso Bilateral"],
            [None] * 12 + ["5", "Blindaje"],
        ],
    )

    webhook = libro.create_sheet("4 - WebHookTransaction")
    for fila, texto in enumerate(["Ambiente URL", "Ambiente DEV", "Ambiente QA", "Ambiente PDN"], start=1):
        webhook.cell(fila, 1, texto)
    _filas(
        webhook,
        5,
        [
            ["Header"],
            ["Authorization: Bearer access_token"],
            ["Metodo"],
            ["Post"],
            [],
            ["Parametros"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["TransactionFlit", "Id del trámite", "number", "S"],
            ["statusMessage", "Mensaje del estado", "String(250)", "N"],
            ["Ejemplo"],
            ['{"transactionFlit": 10, "statusMessage": "Ver ' + RUTA_LOCAL.replace("\\", "\\\\") + '",}'],
            ["Salida"],
            ["Campo", "Descripción", "Tipo de Dato", "Obligatoriedad"],
            ["Status", "Estado de recepción", "Integer"],
            ["Ejemplo"],
            ['{"Status": 1}'],
        ],
    )
    libro.save(ruta)
    return ruta


def _respuesta(ref: str | None = None) -> dict:
    return {"description": "OK", "content": {"application/json": {"schema": {"$ref": ref} if ref else {}}}}


def _salud() -> dict:
    return {
        "get": {
            "operationId": "HealthCheck",
            "responses": {
                "200": {
                    "description": "OK",
                    "headers": {"DefaultResponse": {"schema": {"$ref": "#/components/schemas/DefaultResponse"}}},
                }
            },
            "tags": ["Health: Estado del servicio"],
            "security": [],
            "parameters": [],
        }
    }


def _swagger(paths: dict, schemas: dict, bearer: bool = False) -> dict:
    componentes = {
        "schemas": {
            "DefaultResponse": {
                "type": "object",
                "properties": {"message": {"type": "string"}},
                "required": ["message"],
                "additionalProperties": False,
            },
            **schemas,
        },
        "securitySchemes": {"bearer": {"type": "http", "scheme": "bearer", "bearerFormat": "JWT"}} if bearer else {},
        "requestBodies": {},
        "examples": {},
    }
    return {
        "openapi": "3.0.0",
        "info": {
            "title": "api_express_typescript_boilerplate",
            "version": "1.0.0",
            "description": "Boilerplate. Notas en " + RUTA_LOCAL,
            "contact": {"name": CORREO_AJENO},
        },
        "servers": [{"url": "/"}],
        "paths": {"/api/v1/health": _salud(), **paths},
        "components": componentes,
    }


def crear_repos(carpeta: Path) -> Path:
    tsoa = {"entryFile": "src/index.ts", "spec": {"outputDirectory": "./", "specVersion": 3}}
    swaggers = {
        "BackApiExternalAuth": _swagger(
            {
                "/api/v1/auth/login": {
                    "post": {
                        "operationId": "LoginService",
                        "responses": {"200": _respuesta(), "404": _respuesta("#/components/schemas/DefaultResponse")},
                        "tags": ["Autenticación"],
                        "security": [],
                        "parameters": [],
                        "requestBody": {
                            "required": True,
                            "content": {
                                "application/json": {"schema": {"$ref": "#/components/schemas/LoginServiceRequestDto"}}
                            },
                        },
                    }
                },
                "/api/v1/company-parameters": {
                    "post": {
                        "operationId": "CreateNew",
                        "responses": {"200": _respuesta("#/components/schemas/DefaultResponse")},
                        "tags": ["Clientes"],
                        "security": [],
                        "parameters": [],
                    }
                },
            },
            {
                "LoginServiceRequestDto": {
                    "type": "object",
                    "properties": {"username": {"type": "string"}, "password": {"type": "string"}},
                    "required": ["username", "password"],
                    "additionalProperties": False,
                }
            },
        ),
        "BackApiExternalTransact": _swagger(
            {
                "/api/v1/external-transaction/register": {
                    "post": {
                        "operationId": "CreateNewList",
                        "responses": {"200": _respuesta("#/components/schemas/CreateListResponseModel")},
                        "tags": ["External Transaction"],
                        "security": [{"bearer": []}],
                        "parameters": [],
                        "requestBody": {
                            "required": True,
                            "content": {
                                "application/json": {
                                    "schema": {
                                        "type": "array",
                                        "items": {"$ref": "#/components/schemas/ExternalIntegrationMasterRequestDto"},
                                    }
                                }
                            },
                        },
                    }
                },
                "/api/v1/source-query": {
                    "post": {
                        "operationId": "CreateNew",
                        "responses": {"200": _respuesta("#/components/schemas/DefaultResponse")},
                        "tags": ["Source Query"],
                        "security": [],
                        "parameters": [],
                    }
                },
                "/api/v1/external-transaction": {
                    "post": {
                        "operationId": "CreateNew",
                        "responses": {"200": _respuesta("#/components/schemas/DefaultResponse")},
                        "tags": ["External Transaction"],
                        "security": [{"bearer": []}],
                        "parameters": [],
                    }
                },
            },
            {
                "ExternalIntegrationMasterRequestDto": {
                    "type": "object",
                    "properties": {"companyManagerId": {"type": "number"}, "actors": {"type": "array", "items": {}}},
                },
                "CreateListResponseModel": {
                    "type": "object",
                    "properties": {
                        "TotalRows": {"type": "number", "format": "double"},
                        "Detail": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "Plate": {"type": "string"},
                                    "Status": {"type": "number"},
                                    "Message": {"type": "string"},
                                    "TransactionFlit": {"type": "number", "nullable": True},
                                },
                            },
                        },
                    },
                    "required": ["TotalRows", "Detail"],
                },
            },
            bearer=True,
        ),
        "BackApiExternalTransactWebHook": _swagger({}, {}, bearer=True),
    }
    for repo, swagger in swaggers.items():
        destino = carpeta / repo
        destino.mkdir(parents=True)
        (destino / "swagger.json").write_text(json.dumps(swagger), encoding="utf-8")
        (destino / "tsoa.json").write_text(json.dumps(tsoa), encoding="utf-8")
    return carpeta


@pytest.fixture
def excel(tmp_path: Path) -> Path:
    return crear_excel(tmp_path / "Contratos_Integraciones_Prueba.xlsx")


@pytest.fixture
def repos(tmp_path: Path) -> Path:
    return crear_repos(tmp_path / "repos")
