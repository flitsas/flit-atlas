"""CA de la HU 4.5: el paquete ICT y el de Atlas generan los 7 formatos y cada uno abre y trae lo esperado."""

import io
import json
import re
import zipfile
from pathlib import Path

import pytest
import yaml
from apoyo_exportes import MARCA, ORGANIZACION, requiere_pdf
from docx import Document
from jsonschema import Draft4Validator
from openpyxl import load_workbook

from atlas_docs.exportes import Formato, ResultadoExportacion

ESQUEMA_POSTMAN = Path(__file__).parent / "esquemas" / "postman-collection-v2.1.0.json"


def _uno(resultado: ResultadoExportacion, formato: Formato) -> bytes:
    archivos = resultado.por_formato(formato)
    assert len(archivos) == 1, [a.ruta for a in archivos]
    return archivos[0].datos


# 7 formatos -------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("nombre", ["ict", "atlas"])
def test_genera_los_7_formatos_y_el_zip_los_trae_todos(nombre: str, request: pytest.FixtureRequest) -> None:
    requiere_pdf()
    resultado: ResultadoExportacion = request.getfixturevalue(nombre)

    assert {a.formato for a in resultado.archivos} == set(Formato)
    zip_ = resultado.zip
    assert zip_ is not None
    with zipfile.ZipFile(io.BytesIO(zip_.datos)) as archivo:
        nombres = archivo.namelist()
        base = nombres[0].split("/")[0]
        manifiesto = json.loads(archivo.read(f"{base}/manifiesto.json"))
        leeme = archivo.read(f"{base}/LEEME.md").decode()
        assert archivo.testzip() is None
    esperados = {f"{base}/{a.ruta}" for a in resultado.archivos if a.formato != Formato.ZIP}
    assert esperados | {f"{base}/LEEME.md", f"{base}/manifiesto.json"} == set(nombres)
    assert {a["formato"] for a in manifiesto["archivos"]} == set(Formato) - {Formato.ZIP}
    assert manifiesto["hash_contenido"] == resultado.hash_contenido
    for entrada in manifiesto["archivos"]:
        assert entrada["sha256"] == resultado.archivo(entrada["ruta"]).sha256
    assert resultado.hash_contenido in leeme and ORGANIZACION in leeme


def test_sin_pdf_genera_los_demas_formatos(ict: ResultadoExportacion) -> None:
    formatos = {a.formato for a in ict.archivos}
    assert formatos >= set(Formato) - {Formato.PDF}


# OpenAPI ----------------------------------------------------------------------------------------------------------
def test_openapi_yaml_y_json_dicen_lo_mismo(ict: ResultadoExportacion) -> None:
    yaml_ = next(a for a in ict.por_formato(Formato.OPENAPI) if a.ruta.endswith(".yaml"))
    json_ = next(a for a in ict.por_formato(Formato.OPENAPI) if a.ruta.endswith(".json"))

    documento = yaml.safe_load(yaml_.datos)
    assert documento == json.loads(json_.datos)
    assert documento["openapi"].startswith("3.1")
    # Las fechas siguen siendo texto (no se convierten a fecha al leer ni al escribir).
    assert documento["info"]["x-flit-changelog"][0]["fecha"] == "2025-05-16"


# Excel ------------------------------------------------------------------------------------------------------------
def test_excel_abre_y_tiene_las_hojas_y_columnas_del_formato_actual(ict: ResultadoExportacion) -> None:
    libro = load_workbook(io.BytesIO(_uno(ict, Formato.EXCEL)))

    hojas = libro.sheetnames
    assert hojas[:3] == ["Portada", "Índice", "Cambios"]
    assert "Tablas de referencia" in hojas and "Guías" in hojas
    servicios = [h for h in hojas if re.match(r"^\d+ - ", h)]
    assert len(servicios) == 13  # 12 operaciones y el webhook
    assert all(len(h) <= 31 for h in hojas)

    indice = libro["Índice"]
    assert indice["A1"].value == "Versión 3"
    assert indice["A2"].value == "Fecha actualización: 24/10/2026"
    assert indice["A3"].value == "Listado de servicios integrados"
    assert indice["A4"].hyperlink.location or indice["A4"].hyperlink.target

    registro = libro[next(h for h in servicios if "postExternalTransaction" in h)]
    assert registro["A1"].value == "Ambiente URL"
    assert registro["A1"].fill.fgColor.rgb == "FF4E00EB"
    etiquetas = [c.value for c in registro["A"] if c.value]
    for esperado in ("Ambiente QA", "Ambiente PDN", "Header", "Método", "Parámetros", "Ejemplo"):
        assert esperado in etiquetas
    qa = next(r for r in registro.iter_rows(max_col=2) if r[0].value == "Ambiente QA")
    assert qa[1].value == "https://ict.qa.flitsas.com"
    fila_campos = next(r for r in registro.iter_rows() if r[0].value == "Campo")
    columnas = [c.value for c in fila_campos[:8]]  # a la derecha siguen las tablas de referencia
    assert columnas == [
        "Campo",
        "Descripción",
        "OBLIGATORIO TRASPASO BILATERAL",
        "OBLIGATORIO TRASPASO UNILATERAL",
        "OBLIGATORIO MATRICULA INICIAL",
        "OBLIGATORIO MATRICULA TIPO LEASING",
        "Tipo de Dato",
        "Observaciones",
    ]
    valores = {str(c.value) for fila in registro.iter_rows() for c in fila if c.value is not None}
    assert {"seller/document_type", "Tabla Tipo Documento", "Tabla Tipo Tramite", "S", "N"} <= valores

    adjuntos = libro[next(h for h in servicios if "Attachments" in h)]
    assert "Tabla guía documentos adjuntos" in {c.value for fila in adjuntos.iter_rows() for c in fila}
    assert ict.hash_contenido in {c.value for fila in libro["Portada"].iter_rows() for c in fila}
    assert libro["Portada"].oddHeader.right.text == MARCA


def test_excel_del_contrato_de_atlas_abre(atlas: ResultadoExportacion) -> None:
    libro = load_workbook(io.BytesIO(_uno(atlas, Formato.EXCEL)), read_only=True)
    assert len([h for h in libro.sheetnames if re.match(r"^\d+ - ", h)]) > 100


# Word -------------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("nombre", ["ict", "atlas"])
def test_word_abre_con_portada_marca_de_agua_y_hash(nombre: str, request: pytest.FixtureRequest) -> None:
    resultado: ResultadoExportacion = request.getfixturevalue(nombre)
    documento = Document(io.BytesIO(_uno(resultado, Formato.WORD)))

    textos = [p.text for p in documento.paragraphs]
    assert "Integración ICT" in textos
    assert any(resultado.hash_contenido in t for t in textos)
    assert "Índice" in textos and "Cambios del paquete" in textos
    assert documento.sections[0].header.paragraphs[0].text == MARCA
    assert documento.core_properties.author == "Flit Atlas"
    assert len(documento.tables) > 3


# PDF --------------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("nombre", ["ict", "atlas"])
def test_pdf_tiene_marca_de_agua_hash_y_varias_paginas(nombre: str, request: pytest.FixtureRequest) -> None:
    requiere_pdf()
    from pypdf import PdfReader  # noqa: PLC0415

    resultado: ResultadoExportacion = request.getfixturevalue(nombre)
    lector = PdfReader(io.BytesIO(_uno(resultado, Formato.PDF)))

    assert len(lector.pages) > 1
    portada = lector.pages[0].extract_text()
    assert resultado.hash_contenido in portada.replace("\n", "")
    for pagina in (lector.pages[1], lector.pages[-1]):
        assert MARCA in pagina.extract_text()


# HTML -------------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("nombre", ["ict", "atlas"])
def test_html_es_autocontenido_y_sin_scripts(nombre: str, request: pytest.FixtureRequest) -> None:
    resultado: ResultadoExportacion = request.getfixturevalue(nombre)
    html = _uno(resultado, Formato.HTML).decode()

    assert "<script" not in html.lower()  # ni scripts propios ni el de la guía (sale escapado)
    assert "&lt;script&gt;" in html
    # Ningún atributo carga o enlaza algo fuera del documento.
    for atributo in re.findall(r"""\b(?:src|href|action|poster|data)\s*=\s*["']?([^"'\s>]*)""", html, re.I):
        assert atributo.startswith("#"), atributo
    assert not re.search(r"url\(|@import", html, re.I)
    assert "default-src 'none'" in html
    assert MARCA in html and resultado.hash_contenido in html
    # El enlace externo de la guía queda escrito como texto y la imagen, como su texto alternativo.
    assert "(https://ejemplo.com/guia)" in html and "[imagen: diagrama]" in html


# Postman ----------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("nombre", ["ict", "atlas"])
def test_postman_valida_contra_el_esquema_oficial_v21(nombre: str, request: pytest.FixtureRequest) -> None:
    resultado: ResultadoExportacion = request.getfixturevalue(nombre)
    coleccion = json.loads(_uno(resultado, Formato.POSTMAN))
    esquema = json.loads(ESQUEMA_POSTMAN.read_text(encoding="utf-8"))

    errores = sorted(Draft4Validator(esquema).iter_errors(coleccion), key=lambda e: list(e.path))
    assert not errores, [f"{list(e.path)}: {e.message}" for e in errores[:5]]
    variables = {v["key"]: v["value"] for v in coleccion["variable"]}
    assert set(variables) == {"baseUrl", "token"}
    assert variables["token"] == ""
    peticiones = [i["request"] for carpeta in coleccion["item"] for i in carpeta["item"]]
    assert peticiones and all(p["url"]["raw"].startswith("{{baseUrl}}") for p in peticiones)
    assert '"event":' not in json.dumps(coleccion)  # sin scripts


def test_postman_ict_usa_bearer_token_y_qa(ict: ResultadoExportacion) -> None:
    coleccion = json.loads(_uno(ict, Formato.POSTMAN))
    variables = {v["key"]: v["value"] for v in coleccion["variable"]}

    assert variables["baseUrl"] == "https://ict.qa.flitsas.com"
    assert coleccion["auth"]["bearer"][0]["value"] == "{{token}}"
    peticiones = {i["name"]: i["request"] for carpeta in coleccion["item"] for i in carpeta["item"]}
    login = next(p for n, p in peticiones.items() if "autenticacion" in n.lower())
    assert login["auth"] == {"type": "noauth"}
    estado = next(p for p in peticiones.values() if "/status-process/byId/" in p["url"]["raw"])
    assert estado["url"]["raw"] == "{{baseUrl}}/api/v1/status-process/byId/:id"
    assert estado["url"]["variable"][0]["key"] == "id"


# Hash -------------------------------------------------------------------------------------------------------------
def test_hash_en_la_portada_de_cada_formato(ict: ResultadoExportacion) -> None:
    assert re.fullmatch(r"[0-9a-f]{64}", ict.hash_contenido)
    assert ict.hash_contenido in _uno(ict, Formato.HTML).decode()
    assert ict.hash_contenido in json.loads(_uno(ict, Formato.POSTMAN))["info"]["description"]
